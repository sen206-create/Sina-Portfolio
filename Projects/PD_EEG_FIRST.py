import mne
import numpy as np
import pandas as pd
from pathlib import Path
from mne.preprocessing import ICA
from mne_icalabel import label_components
from pyprep import NoisyChannels
from autoreject import AutoReject


def preprocess_epoch_present(raw, subject_id):
    detector = NoisyChannels(
        raw.copy(),
        random_state=42
    )

    detector.find_all_bads()
    bads = detector.get_bads()

    raw_marked = raw.copy()
    raw_marked.info["bads"] = sorted(
        set(raw.info["bads"]) | set(bads)
    )

    raw_marked.set_eeg_reference("average", projection=False)

    raw_filtered = raw_marked.copy().filter(l_freq=1, h_freq=40.0)
    raw_for_ica = raw_marked.copy().filter(l_freq=1, h_freq=100.0)

    epochs_for_ica = mne.make_fixed_length_epochs(
        raw_for_ica,
        duration=2,
        overlap=0,
        preload=True,
    )
    ar_before_ica = AutoReject(random_state=42)
    ar_before_ica.fit(epochs_for_ica)
    ica_epoch_reject_log = ar_before_ica.get_reject_log(epochs_for_ica)

    epochs_for_ica_fit = epochs_for_ica[~ica_epoch_reject_log.bad_epochs]

    if len(epochs_for_ica_fit) == 0:
        raise ValueError("No epochs remaining for ICA.")
    ica = ICA(
        n_components=None,
        method='infomax',
        fit_params=dict(extended=True),
        random_state=42,
        max_iter='auto',
    )
    ica.fit(epochs_for_ica_fit, picks="eeg")
    ic_labels = label_components(epochs_for_ica_fit, ica, method="iclabel")
    labels = ic_labels["labels"]
    confidence = ic_labels["y_pred_proba"]
    exclude_idx = [
        idx
        for idx, (label, score) in enumerate(zip(labels, confidence))
        if label not in ["brain", "other"] and score >= 0.90
    ]
    ica.exclude = exclude_idx
    raw_clean = ica.apply(raw_filtered.copy())
    raw_clean.interpolate_bads(reset_bads=True)

    epochs = mne.make_fixed_length_epochs(
        raw_clean,
        duration=2,
        overlap=0,
        preload=True
    )
    ar = AutoReject(random_state=42)
    epochs_clean, reject_log = ar.fit_transform(epochs, return_log=True)

    inter_rates = np.mean(reject_log.labels == 2, axis=0)

    frequent_channels = [
        f"{name}:{rate:.1%}"
        for name, rate in zip(reject_log.ch_names, inter_rates)
        if rate > 0.50
    ]

    # PARAMETERS FOR SPREADSHEET
    n_eeg = len(raw_marked.ch_names)
    bad_eeg = raw_marked.info["bads"].copy()
    removed_components_fraction = len(exclude_idx) / ica.n_components_
    retained_epochs_fraction = len(epochs_clean) / len(epochs)
    retained_seconds = len(epochs_clean) * 2

    review_reasons = []

    if (len(bad_eeg)/n_eeg) > 0.20:
        review_reasons.append(f"More than 20% of EEG channels marked bad")

    if removed_components_fraction > 0.30:
        review_reasons.append(f"More than 30% of ICA components removed")

    if frequent_channels:
        review_reasons.append("Channels interpolated in over 50% of epochs")

    if retained_seconds < 60:
        review_reasons.append("Less than 60 seconds retained")

    if ica.n_iter_ >= ica.max_iter:
        review_reasons.append("ICA reached its iteration limit")

    qc = {
        "Subject": subject_id,
        "Status": "Completed",
        "Bad channels": (
            f"{len(bad_eeg)}: {', '.join(bad_eeg)}"
            if bad_eeg else "0"
        ),
        "ICA components removed": (
            f"{len(exclude_idx)}/{ica.n_components_} "
            f"({removed_components_fraction:.1%})"
        ),
        "Epochs retained": (
            f"{len(epochs_clean)}/{len(epochs)} "
            f"({retained_epochs_fraction:.1%})"
        ),
        "Retained duration": retained_seconds,
        "Frequently interpolated channels": (
            "; ".join(frequent_channels) or "None"
        ),
        "Review needed": "Yes" if review_reasons else "No",
        "Review reasons": "; ".join(review_reasons) or "None",
    }

    return epochs_clean, ica, reject_log, qc


def save_qc_spreadsheet(all_qc, output_path):
    columns = [
        "Subject",
        "Status",
        "Bad channels",
        "ICA components removed",
        "Epochs retained",
        "Retained duration",
        "Frequently interpolated channels",
        "Review needed",
        "Review reasons",
    ]

    table = pd.DataFrame(all_qc, columns=columns)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        table.to_excel(
            writer,
            sheet_name="EEG quality control",
            index=False,
        )

        sheet = writer.sheets["EEG quality control"]
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.comments import Comment

        # Formatting
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True)

        sheet.row_dimensions[1].height = 32
        sheet["F1"].comment = Comment(
            "Retained duration is measured in seconds.", "QC"
        )

        widths = {
            "A": 14, "B": 14, "C": 45,
            "D": 25, "E": 25, "F": 20,
            "G": 55, "H": 18, "I": 65,
        }

        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            if row[7].value == "Yes":
                row[7].fill = PatternFill(
                    "solid", fgColor="FFF2CC"
                )

    print(f"Spreadsheet saved to: {output_path}")


idx = [f"{number:03d}" for number in range(1, 2)]

all_qc = []

for index in idx:
    subject_id = f"sub-{index}"

    file_path = (
        f"/Users/sinaenayati/Downloads/v1/"
        f"{subject_id}/eeg/{subject_id}_task-Rest_eeg.set"
    )

    raw = mne.io.read_raw_eeglab(
        file_path,
        preload=True,
    )

    epochs_clean, ica, reject_log, qc = preprocess_epoch_present(
        raw,
        subject_id=subject_id,
    )

    all_qc.append(qc)

    save_qc_spreadsheet(
        all_qc,
        "/Users/sinaenayati/Desktop/Project_Intern/EEG_quality_control.xlsx",
    )
